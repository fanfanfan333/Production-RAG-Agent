"""造一份内嵌深色流程图的 DOCX，并跑一遍入库管线（容器内运行）."""
import asyncio
import io
import sys

from docx import Document
from docx.shared import Inches


def build_docx(image_path: str) -> bytes:
    doc = Document()
    doc.add_heading("深度学习模型实战文档", level=1)
    doc.add_paragraph("项目名称：基于OpenCV+Tesseract的高效OCR文本识别工具")
    doc.add_paragraph("import cv2")
    doc.add_paragraph("//通过卷积，池化，摊平，线性变换，将数据压缩成10类")
    doc.add_paragraph('writer = SummaryWriter("../logs_seq")')
    doc.add_paragraph("writer.add_graph(fan, input)")
    doc.add_paragraph("//查看数据流")
    doc.add_picture(image_path, width=Inches(2.2))
    doc.add_paragraph("上图为模型的数据流图。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def main(image_path: str) -> None:
    from sqlalchemy import select

    from app.db.postgres import get_db_session
    from app.db.user_models import User
    from app.services.document_service import process_uploads

    async with get_db_session() as s:
        admin = (await s.execute(select(User).where(User.role == "admin").limit(1))).scalar_one()
    print("owner:", admin.username, admin.id, "tenant:", admin.tenant_id)

    resp = await process_uploads(
        [("vision_flow_demo.docx", build_docx(image_path))],
        owner_id=admin.id,
        tenant_id=admin.tenant_id or "default",
        department_id=admin.department_id,
        access_level="tenant",
    )
    for r in resp.documents:
        print("ingest:", r.status, r.filename, "chunks:", getattr(r, "chunk_count", None),
              "images:", getattr(r, "image_count", None),
              "image_objects:", getattr(r, "image_object_count", None),
              "doc_id:", r.document_id)


asyncio.run(main(sys.argv[1]))
