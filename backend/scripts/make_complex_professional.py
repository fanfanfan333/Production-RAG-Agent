"""Create 120 grounded acceptance cases with merged DOCX tables, scanned PDFs and charts."""
import argparse
import csv
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from scripts.make_professional_golden import make_golden, verify


def make_complex_golden():
    golden = make_golden()
    golden["name"] = "professional-complex-v2"
    replacements = {f"table-{i:02d}.md": f"table-{i:02d}.{'docx' if i < 2 else 'pdf'}" for i in range(5)}
    golden["documents"] = {replacements.get(name, name): text for name, text in golden["documents"].items()}
    for i in range(5, 15):
        amount = 500 + i * 31
        golden["documents"][f"image-{i:02d}.png"] = f"图表{i + 1} 项目VIZ-{i + 100}收入（万元）\n2023年 {amount - 120}\n2024年 {amount - 60}\n2025年 {amount}"
    for case in golden["cases"]:
        case["expect_sources"] = [replacements.get(name, name) for name in case["expect_sources"]]
        if case["category"] == "image":
            index = int(case["id"].split("-")[-1])
            case["query"] = f"图表{index}中2025年的项目收入是多少？"
        case["source_locations"] = [{"filename": name, "page_number": 1, "evidence": golden["documents"][name]} for name in case["expect_sources"]]
    verify(golden)
    return golden


def write_documents(golden, directory):
    from PIL import Image, ImageDraw, ImageFont
    from docx import Document
    font_path = next((path for path in ("C:/Windows/Fonts/msyh.ttc", "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc") if Path(path).exists()), None)
    if not font_path: raise RuntimeError("需要中文字体，不能生成缺字的测试文档")
    font = ImageFont.truetype(font_path, 38)
    small = ImageFont.truetype(font_path, 30)
    directory.mkdir(parents=True, exist_ok=True)
    for name, text in golden["documents"].items():
        path = directory / name
        if name.endswith(".docx"):
            i = int(name.split("-")[1].split(".")[0])
            doc = Document(); doc.add_heading(f"202{i}预算", 0)
            table = doc.add_table(rows=7, cols=3); table.style = "Table Grid"
            table.cell(0, 0).merge(table.cell(1, 0)).text = "部门"
            table.cell(0, 1).merge(table.cell(0, 2)).text = "预算"
            table.cell(1, 1).text = "金额"; table.cell(1, 2).text = "单位"
            for j in range(5):
                for column, value in enumerate((f"部门{j + 1}", str(100 + i * 100 + j * 7), "万元")):
                    table.cell(j + 2, column).text = value
            doc.save(path)
        elif name.endswith(".pdf"):
            i = int(name.split("-")[1].split(".")[0])
            image = Image.new("RGB", (1500, 1200), "white"); draw = ImageDraw.Draw(image)
            draw.text((120, 100), f"202{i}预算 — 扫描表格", font=font, fill="black")
            left, top, width, height = 120, 260, 1260, 105
            draw.rectangle((left, top, left + width, top + height * 7), outline="black", width=3)
            for x in (left + 420, left + 840): draw.line((x, top, x, top + height * 7), fill="black", width=3)
            for row in range(1, 7): draw.line((left, top + height * row, left + width, top + height * row), fill="black", width=3)
            for column, label in enumerate(("部门", "预算金额", "单位")): draw.text((left + 30 + column * 420, top + 30), label, font=font, fill="black")
            for j in range(5):
                for column, value in enumerate((f"部门{j + 1}", str(100 + i * 100 + j * 7), "万元")):
                    draw.text((left + 30 + column * 420, top + height * (j + 1) + 30), value, font=font, fill="black")
            image.save(path, "PDF", resolution=150)
        elif name.endswith(".png"):
            i = int(name.split("-")[1].split(".")[0])
            image = Image.new("RGB", (1400, 1000), "white"); draw = ImageDraw.Draw(image)
            if i < 5:
                draw.multiline_text((70, 100), text, font=font, fill="black", spacing=30)
            else:
                amount = 500 + i * 31; values = [amount - 120, amount - 60, amount]
                draw.text((80, 70), f"图表{i + 1} 项目VIZ-{i + 100}收入（万元）", font=font, fill="black")
                draw.line((140, 220, 140, 800, 1250, 800), fill="black", width=4)
                maximum = amount + 100
                points = [(350 + j * 350, 800 - int(value / maximum * 500)) for j, value in enumerate(values)]
                for j, ((x, y), value) in enumerate(zip(points, values)):
                    if i < 10: draw.rectangle((x - 70, y, x + 70, 799), fill=(55, 110, 190))
                    else: draw.ellipse((x - 9, y - 9, x + 9, y + 9), fill=(55, 110, 190))
                    draw.text((x - 60, y - 65), str(value), font=font, fill="black")
                    draw.text((x - 70, 835), f"202{j + 3}", font=small, fill="black")
                if i >= 10: draw.line(points, fill=(55, 110, 190), width=5)
                draw.text((70, 910), "数据标签为对应年份的收入；单位：万元。", font=small, fill="black")
            image.save(path)
        else:
            path.write_text(text, encoding="utf-8")
    assert all((directory / name).is_file() for name in golden["documents"])


def attach_public_excerpt(golden, pdf_path, directory):
    import fitz
    import hashlib
    name = "graphrag-public-excerpt.pdf"
    source = fitz.open(pdf_path); excerpt = fitz.open()
    assert len(source) >= 9 and "GraphRAG" in source[0].get_text() and "2404.16130v2" in source[0].get_text(), "需要指定版本的 GraphRAG 原始论文"
    for page in (6, 7, 8): excerpt.insert_pdf(source, from_page=page, to_page=page)
    excerpt.save(directory / name)
    pages = [page.get_text() for page in excerpt]
    references = [
        (1, "GraphRAG论文每个数据集生成了多少个测试问题？", "125"),
        (2, "GraphRAG论文Podcast语料被切分成多少个文本块？", "1669"),
        (2, "GraphRAG论文News语料被切分成多少个文本块？", "3197"),
        (2, "GraphRAG论文实验中每个文本块长度是多少token？", "600"),
        (2, "GraphRAG论文实验中相邻文本块重叠多少token？", "100"),
        (3, "GraphRAG论文Podcast语料的图索引构建耗时多少分钟？", "281"),
    ]
    for case, (page, query, answer) in zip([case for case in golden["cases"] if case["category"] == "fact"][-6:], references):
        assert answer in pages[page - 1], "公开参考答案必须出现在指定原文页面"
        case.update(query=query, answer=answer, expect_sources=[name], review="original_source_quote_verified",
            source_locations=[dict(filename=name, page_number=page, evidence=answer,
                source_url="https://arxiv.org/pdf/2404.16130v2", original_page=page + 6)])
    used = {name for case in golden["cases"] for name in case["expect_sources"]}
    golden["documents"] = {name: text for name, text in golden["documents"].items() if name in used}
    golden["documents"][name] = " ".join(answer for _, _, answer in references)
    golden["public_source"] = dict(url="https://arxiv.org/pdf/2404.16130v2", sha256=hashlib.sha256(pdf_path.read_bytes()).hexdigest(),
        original_pages=[7, 8, 9], use="evaluation only; not training data", human_review="pending")
    verify(golden)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--public-pdf", type=Path, help="Downloaded GraphRAG arXiv 2404.16130v2; only pages 7-9 used")
    args = parser.parse_args()
    golden = make_complex_golden(); write_documents(golden, args.out / "documents")
    if args.public_pdf: attach_public_excerpt(golden, args.public_pdf, args.out / "documents")
    (args.out / "gold.json").write_text(json.dumps(golden, ensure_ascii=False, indent=2), encoding="utf-8")
    with (args.out / "review.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream); writer.writerow(("ID", "类别", "问题", "参考答案", "来源及证据", "人工审核状态", "审核人", "备注"))
        for case in golden["cases"]:
            writer.writerow((case["id"], case["category"], case["query"], case["answer"], json.dumps(case["source_locations"], ensure_ascii=False), "待审核", "", ""))
    print(f"Generated {len(golden['cases'])} cases and {len(golden['documents'])} documents; human review pending")


if __name__ == "__main__": main()
