"""
图片链路端到端验证（容器内运行）.

跑通"上传 → 抽取 → 理解 → 入库 → 检索 → 引用元数据"整条链路，用一份
**真实生成**的带图 DOCX，而不是 mock：

    1. 造一张写着英文文字的图片（Tesseract 只有 eng，中文 OCR 不可用）；
    2. 嵌进 DOCX，与一段正文一起上传；
    3. 走真实的 process_uploads（解析 → 切块 → 向量化 → 入库）；
    4. 用"图里的词"提问，验证图片 chunk **真的能被检索到**；
    5. 检查检索结果携带的图片元数据（image_path / position / content_type），
       这是"图片引用回溯"能否成立的前提。

运行：
    docker cp scripts/e2e_image_chain.py rag_backend:/tmp/
    docker exec rag_backend python /tmp/e2e_image_chain.py
"""

from __future__ import annotations

import asyncio
import io
import sys

from PIL import Image, ImageDraw, ImageFont


# ── 1. 造一张"图里有字"的图片 ────────────────────────────────────────────────
# 用大字号白底黑字，确保 Tesseract 能读出来（eng 语言包）。
IMAGE_WORDS = "INVOICE TOTAL 4820 DOLLARS"


def make_text_image() -> bytes:
    img = Image.new("RGB", (900, 240), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype(
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 56
        )
    except OSError:
        font = ImageFont.load_default()
    draw.text((30, 90), IMAGE_WORDS, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


# ── 2. 造带图 DOCX ───────────────────────────────────────────────────────────

def make_docx(image_bytes: bytes) -> bytes:
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    doc.add_heading("Equipment Manual", level=1)
    doc.add_paragraph(
        "The ABX-300 unit rated power is 3200 watts and it ships with a "
        "standard two year warranty."
    )
    doc.add_paragraph("Billing summary is shown in the figure below.")
    doc.add_picture(io.BytesIO(image_bytes), width=Inches(5.5))
    doc.add_paragraph(
        "Contact the service desk for calibration procedures and spare parts."
    )
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── 3. 主流程 ────────────────────────────────────────────────────────────────

async def main() -> int:
    from app.services.document_service import process_uploads
    from app.services.retrieval_service import retrieve_chunks

    image_bytes = make_text_image()
    docx_bytes = make_docx(image_bytes)
    print(f"[1] built DOCX: {len(docx_bytes):,} bytes, image {len(image_bytes):,} bytes")

    results = await process_uploads([("e2e_image_chain.docx", docx_bytes)])
    doc = results.documents[0]
    print(f"[2] upload: status={doc.status} chunks={doc.chunk_count} pages={doc.page_count}")
    print(
        f"    images: extracted={doc.image_count} indexed_as_object={doc.image_object_count}"
    )
    if doc.status != "completed":
        print(f"    FAILED: {doc.error}")
        return 1
    if not doc.image_count:
        print("    FAILED: 文档里的图片没有被抽取出来")
        return 1
    if not doc.image_object_count:
        print("    FAILED: 图片被抽取但没有建成独立检索对象（可能 OCR 为空）")
        return 1

    doc_id = str(doc.document_id)
    print(f"    document_id={doc_id}")

    # ── 4. 用"图里的词"提问：图片 chunk 必须能被检索到 ──────────────────────
    query = "INVOICE TOTAL 4820 DOLLARS"
    chunks = await retrieve_chunks(query=query, top_k=5, owner_id=None)
    print(f"[3] retrieved {len(chunks)} chunk(s) for {query!r}")

    image_hits = [c for c in chunks if c.image_id]
    text_hits = [c for c in chunks if not c.image_id]
    print(f"    image-derived hits: {len(image_hits)} | text hits: {len(text_hits)}")

    if not image_hits:
        print("    FAILED: 图片 chunk 没被检索到 —— 图片未成为可检索对象")
        for c in chunks:
            print(f"      - {c.filename} p{c.page_number} #{c.chunk_index} "
                  f"[{c.content_type}] score={c.score:.3f}")
        return 1

    hit = image_hits[0]
    print("[4] image chunk metadata (图片引用回溯的前提):")
    for field in (
        "document_id", "filename", "page_number", "chunk_index", "content_type",
        "image_id", "image_path", "image_type", "position", "bbox",
        "analyze_engine", "analyze_confidence", "manual_review",
    ):
        print(f"      {field:20s} = {getattr(hit, field, None)}")
    print(f"      location_label      = {hit.location_label()}")

    ocr_ok = "4820" in (hit.text or "") or "4820" in (hit.image_caption or "")
    print(f"      contains '4820'     = {ocr_ok}")

    # ── 5. 纯文本问题也要能命中正文（图文统一检索不互相挤压）────────────────
    text_chunks = await retrieve_chunks(query="ABX-300 rated power watts", top_k=5)
    print(f"[5] text query returned {len(text_chunks)} chunk(s)")
    for c in text_chunks[:3]:
        print(f"      - {c.filename} p{c.page_number} #{c.chunk_index} "
              f"[{c.content_type}] score={c.score:.3f}")

    ok = bool(image_hits)

    # ── 6. 评测模块接真实检索：算出 Recall / MRR（持续监控的落点）────────────
    from app.services.evaluation import (
        EvalCase,
        EvalSet,
        evaluate,
        item_from_chunk,
    )

    async def _retrieve(q: str):
        got = await retrieve_chunks(query=q, top_k=5, owner_id=None)
        return [item_from_chunk(c) for c in got]

    eval_set = EvalSet(
        name="e2e-image-chain",
        cases=(
            # 图片题：金标写"文件名::页::第几张图"（人类可读，不必背 UUID）
            EvalCase(
                query="INVOICE TOTAL 4820 DOLLARS",
                relevant=frozenset({"e2e_image_chain.docx::p1::1"}),
                modality="image",
            ),
            # 文本题
            EvalCase(
                query="ABX-300 rated power watts",
                relevant=frozenset({"e2e_image_chain.docx::p1"}),
                modality="text",
            ),
            # 负例：知识库里不该有答案
            EvalCase(
                query="quantum chromodynamics lattice gauge theory",
                relevant=frozenset({"__none__"}),
                modality="text",
            ),
        ),
    )
    report = await evaluate(_retrieve, eval_set, k_values=(1, 5))
    print("[6] evaluation report (真实检索):")
    for k, v in report.recall.items():
        print(f"      recall@{k:<3d} = {v}")
    print(f"      mrr          = {report.mrr}")
    for modality, entry in report.by_modality.items():
        print(f"      [{modality}] cases={entry['cases']} mrr={entry['mrr']} "
              f"recall@5={entry.get('recall@5')}")

    # 图片与文本两条腿都要能召回，否则"图文统一检索"就是空话
    img_entry = report.by_modality.get("image", {})
    text_entry = report.by_modality.get("text", {})
    if not (img_entry.get("recall@5") or 0) > 0:
        print("    FAILED: 图片用例 Recall@5 为 0")
        ok = False
    if not (text_entry.get("recall@5") or 0) > 0:
        print("    FAILED: 文本用例 Recall@5 为 0")
        ok = False

    print("\nRESULT:", "PASS — 图片可被独立检索并带回原图元数据" if ok else "FAIL")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
