"""Release gate: real Docling must parse without a native-parser fallback."""
import io
import os
import re
from pathlib import Path

import pytest

pytestmark = pytest.mark.skipif(os.getenv("RUN_REAL_VLM") != "1", reason="real model release gate")


def test_docling_pdf_and_native_table():
    import fitz
    from docx import Document
    from app.config import get_settings
    from app.services.structure.providers import DoclingProvider

    pdf = fitz.open()
    for text in ("ABX-300 rated power: 3200 watts.", "Replace the filter every 90 days."):
        pdf.new_page().insert_text((50, 80), text)
    photo = Path(__file__).resolve().parents[2] / "_doc_photo_fixtures/t0_clean_fullpage.png"
    assert photo.exists(), "complex scanned-page fixture is missing"
    pdf.new_page().insert_image(fitz.Rect(30, 30, 570, 440), filename=str(photo))
    provider = DoclingProvider(get_settings())
    result = provider.parse(pdf.tobytes(), "model-gate.pdf")
    assert result is not None, "Docling failed; model cache must be ready before release"
    assert result.page_count == 3
    assert "3200" in result.markdown and "90 days" in result.markdown
    assert "33500" in result.markdown, "the scanned table page was lost"
    assert result.page_for_offset(result.markdown.index("33500")) == 3
    for value in ("2460", "1810"):
        assert value in result.markdown, f"scanned table lost a key cell: {value}"
    pdf.close()

    doc = Document()
    table = doc.add_table(rows=4, cols=3)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "Metric"
    table.cell(0, 1).merge(table.cell(0, 2)).text = "FY2024"
    table.cell(1, 1).text, table.cell(1, 2).text = "Value", "Unit"
    for row, values in ((2, ("Revenue", "33500", "CNY 10k")), (3, ("Profit", "2460", "CNY 10k"))):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    content = io.BytesIO()
    doc.save(content)
    result = provider.parse(content.getvalue(), "native-table.docx")
    assert result is not None
    assert "Revenue" in result.markdown and "33500" in result.markdown
    for metric, value in (("Revenue", "33500"), ("Profit", "2460")):
        assert re.search(rf"\|\s*{metric}\s*\|\s*{value}\s*\|\s*CNY 10k\s*\|", result.markdown), "merged header shifted a data cell or its unit"
