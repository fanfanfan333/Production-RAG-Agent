"""
图片位置元数据（position + bbox）单元测试 —— 细粒度引用里"图片在哪"的基础.

覆盖：
  * ``_normalize_bbox`` —— 非法值 / NaN / 逆序 / 零面积的收敛规则；
  * PPTX  ``_shape_bbox`` —— EMU → 点换算；
  * PPTX  ``_iter_picture_blobs`` —— 组合形状（group）里的图片坐标平移；
  * PPTX  真实 deck —— 用 python-pptx 造两张已知几何位置的图，走一遍
    ``_iter_picture_blobs_for_slide``，确认拿到的就是页面坐标；
  * ``EmbeddedImageRecognizer`` —— position 只对"真的被收下"的图片递增、
    bbox 被规整后写入 ExtractedImage；
  * PDF ``_image_bbox_on_page`` —— 拿不到矩形时老实返回 None（不编造坐标）；
  * DOCX —— 解析器**不传 bbox**（流式排版无页面几何），position 仍然准确；
  * ``image_position_label`` —— "第 N 张图" / "第 N 个表格"的措辞（单一实现点）；
  * ``location_label`` —— 有行号用行号，图片退化为位置标签；
  * ``_position_fields`` —— Qdrant payload 往返 + 脏数据防御。

为什么要单测这些：坐标是"引用卡片画出哪个框"的唯一依据，错了用户会看到
指向错误位置的框，比没有框更糟。因此这里对"降级为 None"的路径覆盖得比
"成功拿到坐标"的路径更密。

宿主机缺依赖时优雅跳过；推荐在 backend 容器内运行：
    docker exec rag_backend python tests/test_image_position.py
"""

from __future__ import annotations

import io
import os
import sys
import tempfile
import types
from pathlib import Path

_BACKEND_ROOT = str(Path(__file__).resolve().parent.parent)
if _BACKEND_ROOT not in sys.path:
    sys.path.insert(0, _BACKEND_ROOT)

# Settings 需要 POSTGRES_PASSWORD；测试环境给个占位值（不真正连库）
os.environ.setdefault("POSTGRES_PASSWORD", "test-placeholder")

# 把落盘目录重定向到临时目录，避免污染 backend/uploads/
_TMP_ROOT = tempfile.mkdtemp(prefix="rag_pos_test_")
os.environ.setdefault("IMAGE_STORAGE_DIR", os.path.join(_TMP_ROOT, "uploads"))
os.environ.setdefault("DOCUMENT_OUTPUT_DIR", os.path.join(_TMP_ROOT, "generated"))

# 桩包：避免触发 app/services/__init__.py 的重依赖（paddle / docling / torch）。
# 被测模块链上没有 `from app.services import X` 这种包级导出依赖。
if "app.services" not in sys.modules:
    _pkg = types.ModuleType("app.services")
    _pkg.__path__ = [str(Path(_BACKEND_ROOT) / "app" / "services")]
    sys.modules["app.services"] = _pkg

try:
    from PIL import Image, ImageDraw

    from app.services.parsers import docx_parser as docx_mod
    from app.services.parsers import image_recognition as recog_mod
    from app.services.parsers import pdf_parser as pdf_mod
    from app.services.parsers import pptx_parser as pptx_mod
    from app.services.parsers.base import ExtractedImage
    from app.services.parsers.image_recognition import EmbeddedImageRecognizer
    from app.services.retrieval_service import (
        RetrievedChunk,
        _position_fields,
        image_position_label,
    )
except ImportError as exc:  # 宿主机缺依赖 → 跳过（容器内已验证）
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from _module_skip import skip_module

    # 不能用 sys.exit()：pytest 在收集阶段导入本模块，抛 SystemExit 会让整个
    # 会话 INTERNALERROR，同目录其它用例全部跑不了。
    skip_module(f"missing dependency ({exc}) — run inside the backend container")

DOC_ID = "11111111-2222-3333-4444-555555555555"


# ── 工具 ─────────────────────────────────────────────────────────────────────

def _png_bytes(width: int = 200, height: int = 140, colour=(200, 30, 30)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (width, height), colour).save(buf, format="PNG")
    return buf.getvalue()


def _fake_understanding(**overrides):
    """一个足够以假乱真的 ImageUnderstanding —— 避免单测真的去跑 OCR / VLM."""
    classification = types.SimpleNamespace(engine="rules", signals={"h_lines": 0})
    defaults = dict(
        image_type="photo",
        route="ocr",
        classification=classification,
        ocr_text="图内文字",
        ocr_engine="tesseract",
        structured_content=None,
        vision_caption=None,
        analyze_engine="ocr",
        confidence=0.9,
        decision="accept",
        manual_review=False,
        # 质检与双通道融合报告（recognize 会把它们写进 ExtractedImage）
        quality={"ok": True, "score": 1.0, "reasons": [], "checks": ["non-empty"]},
        fusion={"strategy": "ocr-first", "chosen": "ocr", "agreed": True},
    )
    defaults.update(overrides)
    return types.SimpleNamespace(**defaults)


# ── 1. _normalize_bbox ───────────────────────────────────────────────────────

def test_normalize_bbox_accepts_valid() -> None:
    norm = recog_mod._normalize_bbox
    assert norm((10.0, 20.0, 110.0, 80.0)) == (10.0, 20.0, 110.0, 80.0)
    # 逆序分量要自动摆正（旋转页会给出 x0 > x1 的矩形）
    assert norm((110.0, 80.0, 10.0, 20.0)) == (10.0, 20.0, 110.0, 80.0)
    # 整数 / 字符串分量都要能收
    assert norm((0, 0, 72, 36)) == (0.0, 0.0, 72.0, 36.0)
    # 小数统一 round(2)，避免浮点噪声写进 payload
    assert norm((1.234567, 2.345678, 100.111, 200.999)) == (
        1.23, 2.35, 100.11, 201.0,
    )
    print("  ok test_normalize_bbox_accepts_valid")


def test_normalize_bbox_rejects_bad() -> None:
    norm = recog_mod._normalize_bbox
    # None / 空 / 分量个数不对
    assert norm(None) is None
    assert norm(()) is None
    assert norm((1.0, 2.0, 3.0)) is None
    assert norm((1.0, 2.0, 3.0, 4.0, 5.0)) is None
    # 不可转 float
    assert norm(("a", "b", "c", "d")) is None
    # NaN（解析器算坐标时除以 0 会产生）
    assert norm((float("nan"), 0.0, 10.0, 10.0)) is None
    # 零面积：x1 == x2 —— 画出来是一条线，不如没有
    assert norm((5.0, 5.0, 5.0, 9.0)) is None
    assert norm((5.0, 5.0, 9.0, 5.0)) is None
    print("  ok test_normalize_bbox_rejects_bad")


# ── 2. PPTX 几何 ─────────────────────────────────────────────────────────────

def test_pptx_shape_bbox_emu_to_points() -> None:
    """1 点 = 12700 EMU；left/top/width/height 都要正确换算."""
    shape = types.SimpleNamespace(
        left=914400,        # 72 pt
        top=457200,         # 36 pt
        width=2286000,      # 180 pt
        height=1524000,     # 120 pt
    )
    assert pptx_mod._shape_bbox(shape) == (72.0, 36.0, 252.0, 156.0)

    # 属性缺失 / 不是数字 → None（老版 python-pptx 的占位符可能没有几何）
    assert pptx_mod._shape_bbox(types.SimpleNamespace()) is None
    assert pptx_mod._shape_bbox(
        types.SimpleNamespace(left=None, top=0, width=0, height=0)
    ) is None
    print("  ok test_pptx_shape_bbox_emu_to_points")


def _fake_picture(left_pt: float, top_pt: float, w_pt: float, h_pt: float):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    return types.SimpleNamespace(
        shape_type=MSO_SHAPE_TYPE.PICTURE,
        left=int(left_pt * 12700),
        top=int(top_pt * 12700),
        width=int(w_pt * 12700),
        height=int(h_pt * 12700),
        image=types.SimpleNamespace(blob=_png_bytes(30, 20)),
    )


def test_pptx_group_offset_shifts_child_bbox() -> None:
    """
    组合形状里的图片：left/top 是**相对组内**的偏移，必须叠加组的页面位置.

    不叠加就会得到一堆挤在页面左上的框 —— 这是最隐蔽的一类坐标错误。
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    child = _fake_picture(10.0, 5.0, 40.0, 30.0)          # 组内坐标
    group = types.SimpleNamespace(
        shape_type=MSO_SHAPE_TYPE.GROUP,
        left=int(100.0 * 12700),
        top=int(200.0 * 12700),
        width=int(300.0 * 12700),
        height=int(300.0 * 12700),
        shapes=[child],
    )

    got = list(pptx_mod._iter_picture_blobs(group))
    assert len(got) == 1
    _blob, bbox = got[0]
    assert bbox == (110.0, 205.0, 150.0, 235.0), bbox

    # 非组合形状（offset 为 0）保持原坐标
    plain = list(pptx_mod._iter_picture_blobs(_fake_picture(7.0, 8.0, 20.0, 10.0)))
    assert plain[0][1] == (7.0, 8.0, 27.0, 18.0)
    print("  ok test_pptx_group_offset_shifts_child_bbox")


def test_pptx_real_deck_geometry() -> None:
    """
    真造一个 PPTX（两张已知 EMU 位置的图），走一遍幻灯片遍历.

    这是对"EMU 换算 + 遍历顺序"最接近真实输入的一次验证。
    """
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])     # 空白版式
    # 第 1 张：(72pt, 36pt)，180×120pt
    slide.shapes.add_picture(
        io.BytesIO(_png_bytes(200, 60, (200, 30, 30))),
        Emu(914400), Emu(457200), Emu(2286000), Emu(1524000),
    )
    # 第 2 张：(300pt, 200pt)，90×60pt
    slide.shapes.add_picture(
        io.BytesIO(_png_bytes(60, 120, (30, 60, 200))),
        Emu(3810000), Emu(2540000), Emu(1143000), Emu(762000),
    )

    found = list(pptx_mod._iter_picture_blobs_for_slide(slide))
    assert len(found) == 2, f"应遍历到 2 张图片，实际 {len(found)}"
    assert found[0][1] == (72.0, 36.0, 252.0, 156.0), found[0][1]
    assert found[1][1] == (300.0, 200.0, 390.0, 260.0), found[1][1]
    print("  ok test_pptx_real_deck_geometry")


# ── 3. EmbeddedImageRecognizer：position 递增 + bbox 落库 ────────────────────

def test_recognizer_assigns_document_position_and_bbox() -> None:
    """
    position 是**文档级**序号（跨页累计），只对真的被收下的图片递增.

    被过滤（太小）或去重（同一张 logo）的图片不占号 —— 否则"第 3 张图"
    会指向一张用户根本看不到的装饰图。
    """
    original = recog_mod._safe_understand
    try:
        recog_mod._safe_understand = lambda *a, **kw: _fake_understanding()

        recognizer = EmbeddedImageRecognizer("deck.pptx", document_id=DOC_ID)
        # 第 1 页两张（第 2 张会被去重：与第 1 张字节相同）
        assert recognizer.recognize(
            Image.open(io.BytesIO(_png_bytes(200, 140))), page_number=1,
            bbox=(10.0, 20.0, 110.0, 90.0),
        )
        assert not recognizer.recognize(
            Image.open(io.BytesIO(_png_bytes(200, 140))), page_number=1,
        ), "同一张图重复出现应被去重"
        # 太小 → 过滤，不占号
        assert not recognizer.recognize(
            Image.open(io.BytesIO(_png_bytes(4, 4))), page_number=1,
        )
        # 第 2 页一张（带逆序 bbox，应被摆正）
        assert recognizer.recognize(
            Image.open(io.BytesIO(_png_bytes(180, 120, (10, 200, 10)))),
            page_number=2, bbox=(200.0, 150.0, 50.0, 30.0),
        )

        assert len(recognizer.images) == 2, [i.image_id for i in recognizer.images]
        first, second = recognizer.images

        assert first.position == 1 and second.position == 2, "position 必须连续且跨页累计"
        assert first.page_number == 1 and second.page_number == 2
        assert first.bbox == (10.0, 20.0, 110.0, 90.0)
        assert second.bbox == (50.0, 30.0, 200.0, 150.0), "逆序 bbox 应被摆正"
        assert second.image_id.endswith("-p2-i1"), second.image_id
    finally:
        recog_mod._safe_understand = original
    print("  ok test_recognizer_assigns_document_position_and_bbox")


def test_page_scan_also_takes_a_position() -> None:
    """整页扫描同样是一张图片，也要占文档级序号（否则后续序号会错位）."""
    recognizer = EmbeddedImageRecognizer("scan.pdf", document_id=DOC_ID)
    recognizer.register_page_scan(
        Image.new("RGB", (600, 800), (255, 255, 255)),
        page_number=1,
        ocr_text="整页扫描的 OCR 文本",
        bbox=(0.0, 0.0, 595.0, 842.0),
    )
    assert len(recognizer.images) == 1
    img = recognizer.images[0]
    assert img.position == 1
    assert img.bbox == (0.0, 0.0, 595.0, 842.0)
    print("  ok test_page_scan_also_takes_a_position")


# ── 4. PDF / DOCX 的降级行为 ─────────────────────────────────────────────────

def test_pdf_bbox_helper_degrades_quietly() -> None:
    """
    拿不到图片矩形时返回 None —— 不编造坐标.

    PyMuPDF 对矢量图 / 旋转裁剪过的页面可能给不出 rects；此时上层降级为
    只显示"第几页 + 第几张图"，这是**诚实**的降级，不是失败。
    """
    class _Page:
        def __init__(self, rects):
            self._rects = rects

        def get_image_rects(self, _xref):
            if isinstance(self._rects, Exception):
                raise self._rects
            return self._rects

    class _Rect:
        x0, y0, x1, y1 = 30.0, 40.0, 230.0, 180.0

    assert pdf_mod._image_bbox_on_page(_Page([_Rect()]), 7) == (30.0, 40.0, 230.0, 180.0)
    assert pdf_mod._image_bbox_on_page(_Page([]), 7) is None
    assert pdf_mod._image_bbox_on_page(_Page(RuntimeError("boom")), 7) is None
    print("  ok test_pdf_bbox_helper_degrades_quietly")


def test_docx_never_passes_bbox() -> None:
    """
    DOCX 是流式排版，OOXML 里没有绝对页面坐标 → 解析器**不传 bbox**.

    这条断言是"不许假装有坐标"的守门人：一旦有人在 docx_parser 里硬凑
    一个 bbox，引用卡片就会画出错误的框。

    同时也守住 ``source_ordinal``：正文里的 ``<!-- image -->`` 占位符要靠
    "这张图在来源枚举里排第几"才能对齐；被过滤掉的图不占 self.images 的名额，
    漏传就会把图片说明标到错误的占位符上。
    """
    captured: list[dict] = []
    original = recog_mod.EmbeddedImageRecognizer.recognize

    def spy(self, img, page_number=1, *, bbox=None, source_ordinal=None):  # noqa: ANN001
        captured.append(
            {"page": page_number, "bbox": bbox, "source_ordinal": source_ordinal}
        )
        return original(
            self,
            img,
            page_number=page_number,
            bbox=bbox,
            source_ordinal=source_ordinal,
        )

    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_paragraph("正文段落。")
    doc.add_picture(io.BytesIO(_png_bytes(120, 90)))
    buf = io.BytesIO()
    doc.save(buf)

    recog_mod.EmbeddedImageRecognizer.recognize = spy
    try:
        result = docx_mod.DocxParser().parse(buf.getvalue(), "t.docx", document_id=DOC_ID)
    finally:
        recog_mod.EmbeddedImageRecognizer.recognize = original

    assert captured, "DOCX 内嵌图片应走到 recognizer"
    assert all(c["bbox"] is None for c in captured), captured
    # 来源序号必须是 1-based 且与枚举顺序一致（第 1 张图 → 1）
    assert [c["source_ordinal"] for c in captured] == list(
        range(1, len(captured) + 1)
    ), captured
    # position 仍然准确（第 1 张图）
    assert result.image_count >= 1
    assert result.images[0].position == 1
    assert result.images[0].bbox is None
    print("  ok test_docx_never_passes_bbox")


# ── 5. 位置标签措辞（单一实现点）──────────────────────────────────────────────

def test_image_position_label_wording() -> None:
    label = image_position_label
    assert label("image", 2) == "第 2 张图"
    assert label("table", 3) == "第 3 个表格"
    # 表格图片虽然 content_type="table"，措辞要跟着 content_type 走
    assert label("table", 1) == "第 1 个表格"
    # 没有序号 → None（宁可不说，也不说错）
    assert label("image", None) is None
    assert label("image", 0) is None
    print("  ok test_image_position_label_wording")


def test_location_label_prefers_lines_then_position() -> None:
    text_chunk = RetrievedChunk(
        document_id=DOC_ID, filename="年报.pdf", page_number=3, chunk_index=0,
        text="正文", score=0.9, line_start=12, line_end=28,
    )
    assert text_chunk.location_label() == "《年报.pdf》，第 3 页，第 12-28 行"

    image_chunk = RetrievedChunk(
        document_id=DOC_ID, filename="年报.pdf", page_number=3, chunk_index=1,
        text="[图片] 图内文字", score=0.9,
        content_type="image", image_id="d-p3-i1", position=2,
        bbox=(120.0, 80.0, 460.0, 320.0),
    )
    # 图片没有行号 → 用"第 2 张图"；bbox 不塞进自然语言句子
    assert image_chunk.location_label() == "《年报.pdf》，第 3 页，第 2 张图"
    assert "120" not in image_chunk.location_label()
    assert image_chunk.position_span == "第 2 张图"
    assert image_chunk.bbox_span == "x 120-460, y 80-320"
    assert image_chunk.is_image_derived is True

    # 表格图片：is_image_derived 仍为 True（它也是图片产出的）
    table_image = RetrievedChunk(
        document_id=DOC_ID, filename="年报.pdf", page_number=1, chunk_index=2,
        text="| a | b |", score=0.9,
        content_type="table", image_id="d-p1-i1", position=3,
    )
    assert table_image.is_image_derived is True
    assert table_image.position_span == "第 3 个表格"

    # 纯文本块没有 position_span
    assert text_chunk.position_span is None
    # 无位置信息的图片：只到页码为止，不硬凑
    bare = RetrievedChunk(
        document_id=DOC_ID, filename="年报.pdf", page_number=2, chunk_index=3,
        text="[图片]", score=0.8, content_type="image", image_id="d-p2-i1",
    )
    assert bare.location_label() == "《年报.pdf》，第 2 页"
    print("  ok test_location_label_prefers_lines_then_position")


# ── 6. Qdrant payload 往返 ───────────────────────────────────────────────────

def test_position_fields_roundtrip() -> None:
    parsed = _position_fields({"position": 4, "bbox": [10.0, 20.0, 110.0, 90.0]})
    assert parsed == {"position": 4, "bbox": (10.0, 20.0, 110.0, 90.0)}

    # 旧索引没有这两个字段 → None（位置标签自动降级）
    assert _position_fields({}) == {"position": None, "bbox": None}

    # position 是字符串也要能收（Qdrant 里可能被写成 str）
    assert _position_fields({"position": "7"})["position"] == 7
    assert _position_fields({"position": "abc"})["position"] is None

    # 半截坐标整体丢弃 —— 画半个框比不画更容易误导人
    for bad in ([1.0, 2.0], [1.0, 2.0, 3.0, "x"], "nope", 42, None):
        assert _position_fields({"bbox": bad})["bbox"] is None, bad
    print("  ok test_position_fields_roundtrip")


def test_extracted_image_searchable_text_unaffected_by_position() -> None:
    """位置字段是**元数据**，不得影响图片的检索文本（否则检索结果会漂）."""
    img = ExtractedImage(
        image_id="d-p1-i1", page_number=1, ocr_text="季度营收 1.2 亿",
        position=1, bbox=(10.0, 20.0, 110.0, 90.0),
    )
    assert img.searchable_text == "季度营收 1.2 亿"
    assert img.position == 1 and img.bbox == (10.0, 20.0, 110.0, 90.0)

    # 默认值向后兼容
    legacy = ExtractedImage(image_id="d-p1-i2", page_number=1, ocr_text="x")
    assert legacy.position is None and legacy.bbox is None
    print("  ok test_extracted_image_searchable_text_unaffected_by_position")


# ── 执行 ─────────────────────────────────────────────────────────────────────

def main() -> None:
    print("图片位置元数据（position + bbox）单元测试")
    test_normalize_bbox_accepts_valid()
    test_normalize_bbox_rejects_bad()
    test_pptx_shape_bbox_emu_to_points()
    test_pptx_group_offset_shifts_child_bbox()
    test_pptx_real_deck_geometry()
    test_recognizer_assigns_document_position_and_bbox()
    test_page_scan_also_takes_a_position()
    test_pdf_bbox_helper_degrades_quietly()
    test_docx_never_passes_bbox()
    test_image_position_label_wording()
    test_location_label_prefers_lines_then_position()
    test_position_fields_roundtrip()
    test_extracted_image_searchable_text_unaffected_by_position()
    print("\nAll image-position tests passed.")


if __name__ == "__main__":
    main()
