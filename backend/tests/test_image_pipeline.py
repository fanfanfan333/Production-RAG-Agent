"""
图片管线单元测试（部分1 / 部分2 / 部分3 / 三层图片处理 / Document Agent）.

覆盖：
  * image_store —— 落盘路径命名（images/page_3_image_1.png）、读写、目录穿越拒绝；
  * ExtractedImage.searchable_text —— 图片检索文本 = 结构化表格 / OCR / vision；
  * chunker.build_image_chunks —— 图片建成独立 chunk、索引与文本块错开、
    无可检索文本的图片不建块、**表格图片产出 content_type="table"**；
  * chunker.detect_content_type —— Markdown 表格识别为 content_type="table"；
  * vector_service.generate_point_id / VectorPoint —— 图文同号也不撞 ID，
    payload 携带 content_type / image_id / image_path / image_caption；
  * document_agent_service.generate_document —— 真实生成 .docx：
    写入正文段落、把 Markdown 表格转成 Word 表格、插入原始图片；
  * 三层图片处理（Image → Classification → 分流）：
      - Picture Classification 四类判定 + 路由（table→Table Parser，
        chart/diagram→Vision，photo→OCR）；
      - Table Structure Recognition —— 框线切网格 + 逐单元格 OCR → Markdown；
      - StructuredContent —— {type,page,content} 形状与 Qdrant payload 映射；
      - 分类关闭时的退化行为、还原失败时的明确失败。

宿主机缺依赖时优雅跳过（与既有单测同一约定）；推荐在 backend 容器内运行：
    docker exec rag_backend python tests/test_image_pipeline.py
"""

from __future__ import annotations

import io
import os
import sys
import tempfile
from pathlib import Path

_BACKEND_ROOT = str(Path(__file__).resolve().parent.parent)
if _BACKEND_ROOT not in sys.path:
    sys.path.insert(0, _BACKEND_ROOT)

# 把落盘目录重定向到临时目录，避免污染 backend/uploads/
_TMP_ROOT = tempfile.mkdtemp(prefix="rag_img_test_")
os.environ.setdefault("IMAGE_STORAGE_DIR", os.path.join(_TMP_ROOT, "uploads"))
os.environ.setdefault("DOCUMENT_OUTPUT_DIR", os.path.join(_TMP_ROOT, "generated"))

try:
    from PIL import Image

    from app.config import get_settings
    from app.services.chunker import (
        build_chunks,
        build_image_chunks,
        contains_markdown_table,
        detect_content_type,
    )
    from app.services.document_agent_service import generate_document
    from app.services.image_understanding import (
        IMAGE_TYPE_CHART,
        IMAGE_TYPE_DIAGRAM,
        IMAGE_TYPE_PHOTO,
        IMAGE_TYPE_SCREENSHOT,
        IMAGE_TYPE_TABLE,
        ImageUnderstanding,
        classify_image_safe,
        compute_signals,
        content_type_for,
        detect_rules,
        recognize_table,
        resolve_route,
    )
    from app.services.image_understanding.structured_content import StructuredContent
    from app.services.parsers.base import ExtractedImage
    from app.services.storage import (
        image_relative_path,
        resolve_image_path,
        save_image,
    )
    from app.services.vector_service import VectorPoint, generate_point_id
except ImportError as exc:  # 宿主机缺依赖 → 跳过（容器内已验证）
    # 关键：**不要**用 sys.exit() 跳过。pytest 在**收集阶段**导入本模块，此时抛
    # SystemExit 会让整个会话 INTERNALERROR 崩掉 —— 同目录其它用例一条都跑不了，
    # 一个"环境缺依赖"的问题被放大成"整套测试不可用"。helper 在 pytest 下抛
    # Skipped，直接 `python tests/xxx.py` 时才安静退出（详见 tests/_module_skip.py）。
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from _module_skip import skip_module

    skip_module(f"missing dependency ({exc}) — run inside the backend container")

get_settings.cache_clear()
_SETTINGS = get_settings()


# ── 工具 ─────────────────────────────────────────────────────────────────────

def _png_bytes(width: int = 120, height: int = 90, colour=(200, 30, 30)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (width, height), colour).save(buf, format="PNG")
    return buf.getvalue()


DOC_ID = "11111111-2222-3333-4444-555555555555"


# ── 1. 落盘与路径 ────────────────────────────────────────────────────────────

def test_image_relative_path() -> None:
    assert image_relative_path(3, 1) == "images/page_3_image_1.png"
    assert image_relative_path(5, 2, "png") == "images/page_5_image_2.png"
    assert image_relative_path(1, 1, "jpeg") == "images/page_1_image_1.jpg"
    print("  ok test_image_relative_path")


def test_save_and_resolve() -> None:
    rel = save_image(DOC_ID, 3, 1, _png_bytes())
    assert rel == "images/page_3_image_1.png", rel

    resolved = resolve_image_path(DOC_ID, rel)
    assert resolved is not None and resolved.is_file(), "落盘文件应可解析"
    assert resolved.read_bytes()[:4] == b"\x89PNG", "应写入了真实 PNG 数据"

    # 目录穿越必须被拒绝
    assert resolve_image_path(DOC_ID, "../../etc/passwd") is None
    assert resolve_image_path(DOC_ID, "/etc/passwd") is None
    # 不存在的文件返回 None
    assert resolve_image_path(DOC_ID, "images/nope.png") is None
    # 空数据不落盘
    assert save_image(DOC_ID, 1, 9, b"") is None
    print("  ok test_save_and_resolve")


# ── 2. 图片检索文本 ──────────────────────────────────────────────────────────

def test_searchable_text() -> None:
    ocr_only = ExtractedImage(image_id="a", page_number=1, ocr_text="季度营收 1.2 亿")
    assert "1.2 亿" in ocr_only.searchable_text

    caption_only = ExtractedImage(
        image_id="b", page_number=2, ocr_text="", vision_caption="一张卷积神经网络结构图"
    )
    assert "卷积神经网络" in caption_only.searchable_text
    assert caption_only.searchable_text.startswith("图片描述:")

    both = ExtractedImage(
        image_id="c", page_number=3, ocr_text="准确率 95%", vision_caption="折线图"
    )
    text = both.searchable_text
    assert "准确率 95%" in text and "折线图" in text

    empty = ExtractedImage(image_id="d", page_number=4)
    assert empty.searchable_text == "", "既无 OCR 也无 caption 时不应有检索文本"
    print("  ok test_searchable_text")


# ── 3. 图片作为独立检索对象 ──────────────────────────────────────────────────

def test_delete_document_images() -> None:
    """删除文档时必须连带清理落盘目录，否则原图/归档会永久泄漏."""
    from app.services.storage import delete_document_images

    doc_id = "99999999-8888-7777-6666-555555555555"
    rel = save_image(doc_id, 1, 1, _png_bytes())
    assert rel is not None
    resolved = resolve_image_path(doc_id, rel)
    assert resolved is not None and resolved.is_file()

    delete_document_images(doc_id)

    assert resolve_image_path(doc_id, rel) is None, "删除后原图不应再可解析"
    # 重复删除必须幂等（目录已不存在也不报错）
    delete_document_images(doc_id)
    print("  ok test_delete_document_images")


def test_build_image_chunks() -> None:
    images = [
        ExtractedImage(image_id="d-p1-i1", page_number=1, ocr_text="图一中的文字"),
        ExtractedImage(image_id="d-p2-i1", page_number=2, ocr_text=""),  # 无文本 → 不建块
        ExtractedImage(
            image_id="d-p3-i1", page_number=3, ocr_text="", vision_caption="流程图"
        ),
    ]
    chunks = build_image_chunks(images, start_index=7)

    assert len(chunks) == 2, f"应只对有文本的图片建块，实际 {len(chunks)}"
    first, second = chunks
    assert first.content_type == "image"
    assert first.image_id == "d-p1-i1" and first.image_path is None
    assert first.chunk_index == 7, "索引应从 start_index 起（与文本块错开）"
    assert second.chunk_index == 8
    assert second.page_number == 3
    assert second.image_caption == "流程图"
    print("  ok test_build_image_chunks")


def test_table_content_type() -> None:
    assert contains_markdown_table("| a | b |\n| --- | --- |\n| 1 | 2 |") is True
    assert contains_markdown_table("普通段落，没有表格。") is False
    assert contains_markdown_table("| 只有一行 | 不算表格 |") is False

    assert detect_content_type("| a | b |\n| --- | --- |\n| 1 | 2 |") == "table"
    assert detect_content_type("普通段落") == "text"
    # 表格类文档（csv/xlsx）整体标记
    assert detect_content_type("a | b\n1 | 2", default="table") == "table"

    # build_chunks 应把含表格的块标记为 table
    md = "## 数据\n\n| 指标 | 值 |\n| --- | --- |\n| 营收 | 100 |\n\n" + "补充说明。" * 40
    chunks = build_chunks(md, min_chunk_size=50, max_chunk_size=2000, chunk_overlap=20)
    assert any(c.content_type == "table" for c in chunks), "含 Markdown 表格的块应为 table"
    print("  ok test_table_content_type")


# ── 4. Qdrant payload（部分3）────────────────────────────────────────────────

def test_point_id_and_payload() -> None:
    text_id = generate_point_id(
        DOC_ID, "a.pdf", 3, None, None, 5, "同名文本", content_type="text"
    )
    image_id_pid = generate_point_id(
        DOC_ID, "a.pdf", 3, None, None, 5, "同名文本", content_type="image", image_id="x"
    )
    assert text_id != image_id_pid, "图文同 index 同文本也必须得到不同 point id"
    # 确定性：同样输入产生同样 id（幂等续传依赖它）
    assert text_id == generate_point_id(
        DOC_ID, "a.pdf", 3, None, None, 5, "同名文本", content_type="text"
    )

    point = VectorPoint(
        vector=[0.1, 0.2],
        document_id=DOC_ID,
        filename="a.pdf",
        chunk_index=9,
        page_number=3,
        text="[图片] 图内文字",
        content_type="image",
        image_id="d-p3-i1",
        image_path="images/page_3_image_1.png",
        image_caption="一张结构图",
    )
    assert point.content_type == "image"
    assert point.image_path == "images/page_3_image_1.png"
    # 默认值保持向后兼容
    legacy = VectorPoint(
        vector=[0.0], document_id=DOC_ID, filename="b.txt",
        chunk_index=0, page_number=1, text="hello",
    )
    assert legacy.content_type == "text"
    assert legacy.image_id is None and legacy.image_path is None
    print("  ok test_point_id_and_payload")


# ── 5. Document Agent：真实生成 .docx ────────────────────────────────────────

class _FakeChunk:
    def __init__(self, **kw):
        self.document_id = DOC_ID
        self.filename = "深度学习.pdf"
        self.page_number = kw.pop("page_number", 1)
        self.chunk_index = kw.pop("chunk_index", 0)
        self.text = kw.pop("text", "")
        self.score = kw.pop("score", 0.9)
        self.content_type = kw.pop("content_type", "text")
        self.heading = kw.pop("heading", None)
        self.parent_text = kw.pop("parent_text", None)
        self.image_id = kw.pop("image_id", None)
        self.image_path = kw.pop("image_path", None)
        self.image_caption = kw.pop("image_caption", None)


def test_document_agent_generates_docx() -> None:
    from docx import Document as DocxDocument

    # 真实落盘一张图片，供 Agent 插入
    rel = save_image(DOC_ID, 7, 1, _png_bytes(160, 120))

    chunks = [
        _FakeChunk(
            page_number=1,
            chunk_index=0,
            heading="第一章 概述",
            text="这是第一章的正文内容。",
            content_type="text",
        ),
        _FakeChunk(
            page_number=2,
            chunk_index=1,
            text="| 指标 | 数值 |\n| --- | --- |\n| 准确率 | 95% |\n| 召回率 | 92% |",
            content_type="table",
        ),
        _FakeChunk(
            page_number=7,
            chunk_index=2,
            text="[图片] 卷积神经网络结构图",
            content_type="image",
            image_id="d-p7-i1",
            image_path=rel,
            image_caption="卷积神经网络结构图",
        ),
    ]

    info = generate_document("生成一份深度学习调研报告", chunks, title="深度学习调研报告")

    assert info.error is None, f"生成不应失败：{info.error}"
    assert info.filename.endswith(".docx")
    assert info.download_url.endswith(info.filename)
    assert info.section_count == 2, f"文本+表格应有 2 个小节，实际 {info.section_count}"
    assert info.table_count == 1, f"应写入 1 个 Word 表格，实际 {info.table_count}"
    assert info.image_count == 1, f"应插入 1 张图片，实际 {info.image_count}"

    docx_path = resolve_generated_path(info.filename)
    assert docx_path is not None and docx_path.is_file()

    # 打开产物校验真实内容
    doc = DocxDocument(str(docx_path))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "深度学习调研报告" in full_text
    assert "第一章 概述" in full_text
    assert "这是第一章的正文内容。" in full_text
    assert "卷积神经网络结构图" in full_text
    assert len(doc.tables) >= 1, "Markdown 表格应转成真正的 Word 表格"
    assert doc.tables[0].rows[0].cells[0].text.strip() == "指标"

    # 图片真实嵌入（docx 内应有 image 关系）
    image_parts = [
        p for p in doc.part.package.iter_parts()
        if p.content_type.startswith("image/")
    ]
    assert image_parts, "生成的 docx 内应包含嵌入图片"

    # 输出目录中的文件不可穿越
    from app.services.document_agent_service import resolve_generated_file
    assert resolve_generated_file("../secrets.docx") is None
    print("  ok test_document_agent_generates_docx")


def resolve_generated_path(filename: str):
    from app.services.document_agent_service import resolve_generated_file
    return resolve_generated_file(filename)


# ── 6. 意图路由：document_agent ──────────────────────────────────────────────

def test_docx_table_and_image_extraction() -> None:
    """DOCX 解析必须同时拿到：正文段落、表格（转 Markdown）、嵌入图片."""
    import io as _io

    from docx import Document as DocxDocument

    from app.services.parsers.docx_parser import DocxParser

    doc = DocxDocument()
    doc.add_heading("评估报告", level=1)
    doc.add_paragraph("下面是评估指标：")
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "指标"
    table.rows[0].cells[1].text = "数值"
    table.rows[1].cells[0].text = "准确率"
    table.rows[1].cells[1].text = "95%"
    table.rows[2].cells[0].text = "召回率"
    table.rows[2].cells[1].text = "92%"
    doc.add_paragraph("结束语。")
    doc.add_picture(_io.BytesIO(_png_bytes(100, 80)))

    buf = _io.BytesIO()
    doc.save(buf)

    result = DocxParser().parse(buf.getvalue(), "t.docx", document_id=DOC_ID)
    text = result.full_text

    # 表格内容必须出现（python-docx 的 .paragraphs 不含表格，是最易踩的坑）
    assert "准确率" in text and "95%" in text, f"表格内容丢失：\n{text}"
    assert "召回率" in text and "92%" in text
    assert "| --- |" in text, "表格应转成 Markdown 形式"
    assert "下面是评估指标：" in text, "正文段落不应丢失"

    # 含 Markdown 表格的文本应被识别为 table 类型（表格检索依赖它）
    assert detect_content_type(text) == "table"

    # 嵌入图片同样要被结构化抽出（部分1）
    assert result.image_count >= 1, "DOCX 内嵌图片应被抽出为结构化对象"
    print("  ok test_docx_table_and_image_extraction")


def test_document_agent_intent() -> None:
    from app.services.routers.intent_rules import deterministic_route

    assert deterministic_route("帮我生成一份调研报告") == "document_agent"
    assert deterministic_route("把要点整理成一份 Word 文档") == "document_agent"
    assert deterministic_route("导出为 word") == "document_agent"

    # 不能误伤"总结"（要的是摘要文字，不是交付文档）
    assert deterministic_route("总结一下这份报告") != "document_agent"
    assert deterministic_route("这份文档主要讲了什么") != "document_agent"
    # 列表 / 关联 / 闲聊仍然各归各位
    assert deterministic_route("知识库里有哪些文档？") == "list_documents"
    assert deterministic_route("你好") == "general_chat"
    print("  ok test_document_agent_intent")


# ── 7. 三层图片处理：分类 / 表格结构识别 / 路由 ──────────────────────────────

def _table_image(rows=None, with_rules: bool = True):
    """造一张（带/不带框线的）表格图片，用于确定性地测试分类与结构还原."""
    from PIL import ImageDraw

    rows = rows or [["Metric", "Value"], ["Accuracy", "95%"], ["Recall", "92%"]]
    img = Image.new("RGB", (600, 60 + 80 * len(rows)), (255, 255, 255))
    draw = ImageDraw.Draw(img)
    xs = [20, 260, 580]
    ys = [20 + 80 * i for i in range(len(rows) + 1)]
    if with_rules:
        for y in ys:
            draw.line([(xs[0], y), (xs[-1], y)], fill=(0, 0, 0), width=3)
        for x in xs:
            draw.line([(x, ys[0]), (x, ys[-1])], fill=(0, 0, 0), width=3)
    for r, row in enumerate(rows):
        for c, cell in enumerate(row[:2]):
            draw.text((xs[c] + 20, ys[r] + 26), cell, fill=(0, 0, 0))
    return img


def _chart_image():
    """柱状图：白底 + 4 块大面积彩色 + 坐标轴."""
    from PIL import ImageDraw

    img = Image.new("RGB", (520, 320), (255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.line([(60, 280), (480, 280)], fill=(60, 60, 60), width=2)
    draw.line([(60, 280), (60, 40)], fill=(60, 60, 60), width=2)
    for i, colour in enumerate(
        [(66, 133, 244), (219, 68, 55), (244, 180, 0), (15, 157, 88)]
    ):
        x0 = 90 + i * 95
        height = [120, 190, 80, 150][i]
        draw.rectangle([x0, 280 - height, x0 + 60, 280], fill=colour)
    return img


def _diagram_image():
    """流程图：纯线稿（白底黑框）+ 少量文字."""
    from PIL import ImageDraw

    img = Image.new("RGB", (520, 240), (255, 255, 255))
    draw = ImageDraw.Draw(img)
    for box in [(20, 90, 120, 150), (200, 90, 300, 150), (380, 90, 480, 150)]:
        draw.rectangle(box, outline=(0, 0, 0), width=3)
    draw.line([(125, 120), (195, 120)], fill=(0, 0, 0), width=3)
    draw.line([(300, 120), (375, 120)], fill=(0, 0, 0), width=3)
    return img


def _photo_image():
    """照片化：连续渐变 + 噪声，无规则线条、无统一背景."""
    import random

    random.seed(11)
    img = Image.new("RGB", (520, 320))
    px = img.load()
    for y in range(320):
        for x in range(520):
            base = int(90 + 70 * (x / 520.0) + 40 * (y / 320.0))
            px[x, y] = (
                min(255, base + random.randint(-18, 18)),
                min(255, int(base * 0.72) + random.randint(-18, 18)),
                min(255, int(base * 0.5) + random.randint(-18, 18)),
            )
    return img


def test_classify_image_types() -> None:
    """
    Picture Classification：四类图片各自判对，且路由到正确的处理路径.

    这是三层图片处理的第一层 —— 判错了后面全错，所以四类都要覆盖。
    """
    expectations = [
        (_table_image(), IMAGE_TYPE_TABLE, "table_parser"),
        (_chart_image(), IMAGE_TYPE_CHART, "vision"),
        (_diagram_image(), IMAGE_TYPE_DIAGRAM, "vision"),
        (_photo_image(), IMAGE_TYPE_PHOTO, "ocr"),
    ]
    for image, expected_type, expected_route in expectations:
        result = classify_image_safe(image, filename=expected_type)
        assert result.image_type == expected_type, (
            f"期望 {expected_type}，实际 {result.image_type} "
            f"(signals={result.signals})"
        )
        assert resolve_route(result.image_type) == expected_route
        assert 0.0 < result.confidence <= 1.0
    print("  ok test_classify_image_types")


def test_dark_node_diagram_not_chart() -> None:
    """
    深色底节点图的信号 → 必须判成 diagram，不能判成 chart.

    回归保护。下面这组数值不是编的，是**真实误判样本**（`实战.docx` 里那张
    TensorBoard 节点图，input→Fan→output）经预处理后实测出来的：

        h_lines=0 v_lines=0   —— 圆角节点框太短，够不上"贯穿全图"的直线阈值
        chromatic_blocks=2    —— 深底上的橄榄/棕色**描边**凑出 2 个色块
        colorful_ratio=0.098  —— 但彩色只是描边，面积远小于真图表

    老规则只要 ``2<=chromatic<=8`` 就判 chart，于是套用"横轴/纵轴/数据点"模板，
    吐出满篇"无具体数值"（用户反馈的"分析不到位"）。修复后 chart 还要求
    "数据色块面积可观（colorful_ratio≥0.15）**或**存在坐标轴线"，节点图两条都
    不满足，落到 diagram 走 nodes/edges 提示词。这里直接喂信号，精确锁住这两条
    阈值，且不受像素/OCR 环境影响。
    """
    from app.services.image_understanding.classifier import _decide

    settings = get_settings()
    zero_layout = {"ocr_rows": 0, "ocr_cols": 0, "ocr_multi_cell_rows": 0, "ocr_alignment": 0.0}
    zero_text = {"code_score": 0.0, "code_language": "text", "math_ratio": 0.0, "avg_line_chars": 0.0}

    node_signals = {
        "width": 203, "height": 239, "h_lines": 0, "v_lines": 0,
        "inverted": True, "polarity_source": "page-ring",
        "text_bands": 7, "line_art_ratio": 0.862, "saturation": 0.0502,
        "flat_blocks": 4, "chromatic_blocks": 2, "dominant_ratio": 0.7535,
        "colorful_ratio": 0.0983, "ocr_rows": 3, "ocr_cols": 2,
        "ocr_multi_cell_rows": 0, "ocr_alignment": 0.0,
        **zero_text,
    }
    node_signals["avg_line_chars"] = 4.67
    image_type, confidence, reason = _decide(node_signals, settings)
    assert image_type == IMAGE_TYPE_DIAGRAM, (
        f"深色底节点图被判成 {image_type}（{reason}）—— chart 会套用数据点模板"
    )
    assert resolve_route(image_type) == "vision"

    # 反向保护：真柱状图的信号仍须判成 chart —— 别把阈值收得过紧、误伤图表。
    chart_signals = {
        "width": 520, "height": 320, "h_lines": 1, "v_lines": 2,
        "inverted": False, "polarity_source": "page-ring",
        "text_bands": 1, "line_art_ratio": 0.7804, "saturation": 0.1618,
        "flat_blocks": 4, "chromatic_blocks": 3, "dominant_ratio": 0.7424,
        "colorful_ratio": 0.2154, **zero_layout, **zero_text,
    }
    chart_type, _, chart_reason = _decide(chart_signals, settings)
    assert chart_type == IMAGE_TYPE_CHART, (
        f"真柱状图被判成 {chart_type}（{chart_reason}）—— 阈值收过头了"
    )
    print("  ok test_dark_node_diagram_not_chart")


def test_classification_disabled_falls_back_to_ocr() -> None:
    """关闭分类开关时必须整体退化为"普通图片 + OCR"，不能报错."""
    import app.services.image_understanding.classifier as clf

    original = clf.get_settings
    try:
        settings = get_settings()
        settings.IMAGE_CLASSIFICATION_ENABLED = False
        clf.get_settings = lambda: settings            # type: ignore[assignment]
        result = clf.classify_image(_table_image())
        assert result.image_type == IMAGE_TYPE_PHOTO
        assert result.engine == "disabled"
    finally:
        clf.get_settings = original                    # type: ignore[assignment]
        settings.IMAGE_CLASSIFICATION_ENABLED = True
    print("  ok test_classification_disabled_falls_back_to_ocr")


def test_table_image_content_type_mapping() -> None:
    """
    表格图片产出的 chunk 必须是 content_type="table"（而不是 "image"）.

    这是"表格检索"能作用于**图片里的表格**的前提：图片表格与正文表格
    在检索层同构。
    """
    assert content_type_for(IMAGE_TYPE_TABLE) == "table"
    for other in (IMAGE_TYPE_CHART, IMAGE_TYPE_DIAGRAM, IMAGE_TYPE_PHOTO):
        assert content_type_for(other) == "image"

    table_img = ExtractedImage(
        image_id="d-p1-i1", page_number=1,
        image_type=IMAGE_TYPE_TABLE,
        structured_content="| 指标 | 数值 |\n|---|---|\n| 准确率 |95%|",
        ocr_text="指标 数值 准确率 95%",
    )
    assert table_img.content_type == "table"
    # 结构化内容优先，且不再叠加 OCR 文本（避免同一份内容被重复计权）
    assert table_img.searchable_text.startswith("| 指标 |")
    assert "指标 数值 准确率" not in table_img.searchable_text

    # 结构还原失败时仍是图片（不能产出 content_type=table 却没有表格内容）
    degraded = ExtractedImage(
        image_id="d-p1-i2", page_number=1,
        image_type=IMAGE_TYPE_TABLE, structured_content=None, ocr_text="表格文字",
    )
    assert degraded.content_type == "image"
    assert degraded.searchable_text == "表格文字"

    chunks = build_image_chunks([table_img, degraded], start_index=3)
    assert chunks[0].content_type == "table" and chunks[0].image_type == "table"
    assert chunks[1].content_type == "image"
    print("  ok test_table_image_content_type_mapping")


def test_structured_content_shape_and_payload() -> None:
    """结构化内容对象 = 设计稿的 {type, page, content}，并能映射成 Qdrant payload."""
    content = StructuredContent(
        type=IMAGE_TYPE_TABLE,
        page=1,
        content="| 指标 | 数值 |\n|---|---|\n| 准确率 |95%|",
        engine="table_parser",
    )
    assert content.to_dict() == {
        "type": "table",
        "page": 1,
        "content": "| 指标 | 数值 |\n|---|---|\n| 准确率 |95%|",
    }
    payload = content.to_qdrant_payload(source="照片.docx")
    assert payload["content_type"] == "table"
    assert payload["page_number"] == 1
    assert payload["source"] == "照片.docx"
    assert payload["text"].startswith("| 指标 |")
    assert StructuredContent(type="photo", page=1, content="  ").is_empty is True
    print("  ok test_structured_content_shape_and_payload")


def test_recognize_table_with_rules_and_cell_ocr() -> None:
    """
    Table Parser 主线：框线切网格 → 逐单元格 OCR → Markdown 表格.

    用一个按调用顺序返回内容的假 OCR，因此不依赖任何 OCR 引擎，
    可确定性地验证"网格切对了、内容填对了、行列没有多出来"。
    """
    rows = [["Metric", "Value"], ["Accuracy", "95%"], ["Recall", "92%"]]
    img = _table_image(rows)

    row_rules, col_rules = detect_rules(img)
    assert len(row_rules) == len(rows) + 1, (
        f"应检测出 {len(rows) + 1} 条横线，实际 {len(row_rules)}"
    )
    assert len(col_rules) == 3, f"应检测出 3 条竖线，实际 {len(col_rules)}"

    flat = iter([cell for row in rows for cell in row])

    def fake_ocr(_cell_image):
        return next(flat, "")

    structure = recognize_table(img, ocr_fn=fake_ocr)
    assert structure.ok, f"表格还原失败：{structure.meta}"
    assert structure.method == "rules+cell-ocr"
    # 只取相邻框线之间的区间 —— 不能凭空多出页边距那一圈空白行列
    assert structure.rows == len(rows), f"行数应为 {len(rows)}，实际 {structure.rows}"
    assert structure.cols == 2, f"列数应为 2，实际 {structure.cols}"
    lines = structure.markdown.splitlines()
    assert lines[0] == "| Metric | Value |"
    assert lines[1] == "|---|---|"
    assert lines[2] == "| Accuracy | 95% |"
    assert lines[3] == "| Recall | 92% |"
    print("  ok test_recognize_table_with_rules_and_cell_ocr")


def test_recognize_table_without_rescue_degrades() -> None:
    """没有 OCR 回调且没有行坐标时，表格还原必须明确失败（而不是产出坏表格）."""
    structure = recognize_table(_table_image(), lines=[], ocr_fn=None)
    assert not structure.ok
    assert structure.markdown == ""
    assert structure.method == "failed"
    print("  ok test_recognize_table_without_rescue_degrades")


def test_understanding_defaults_are_backward_compatible() -> None:
    """ImageUnderstanding 的默认值必须向后兼容（等同旧行为：图片 + OCR）."""
    result = ImageUnderstanding(
        image_type=IMAGE_TYPE_PHOTO,
        route="ocr",
        classification=classify_image_safe(_photo_image()),
    )
    assert result.analyze_engine == "ocr"
    assert result.structured_content is None
    assert result.vision_caption is None
    assert result.degraded is False
    payload = result.to_dict()
    assert payload["route"] == "ocr" and payload["image_type"] == "photo"
    print("  ok test_understanding_defaults_are_backward_compatible")


def test_signals_are_serialisable() -> None:
    """分类 signals 会被写进日志与 DB，必须是可 JSON 序列化的纯量."""
    import json

    signals = compute_signals(_chart_image())
    json.dumps(signals)      # 不可序列化会直接抛错
    assert "inverted" in signals and "text_bands" in signals
    assert isinstance(signals["h_lines"], int)
    print("  ok test_signals_are_serialisable")


# ── 执行 ─────────────────────────────────────────────────────────────────────

def _blank_signals(**overrides) -> dict:
    """
    构造"形状完整、取值中性"的判定信号，供直接调用 ``_decide`` 的用例使用.

    为什么不手工拼字典：``_decide`` 会读取十来个像素信号（text_bands /
    chromatic_blocks / colorful_ratio / dominant_ratio / saturation /
    line_art_ratio …），它们平时由 ``compute_signals`` 从像素算出来。手工拼
    就必须逐个跟上这些字段，**漏一个会在下游分支炸出 KeyError** —— 报错位置
    离真正的原因（夹具不完整）很远，排查代价极高（本文件踩过这个坑）。

    这里改用一张纯白画布走**真实**计算：形状永远和生产路径一致，``_decide``
    将来新增信号也不用回来改测试。纯白 = 无框线、无彩色块、无文字带，所有
    像素类分支都不命中，于是判定结果只由 ``overrides`` 里的文本类信号决定
    —— 正是公式护栏要单独验证的那一个变量。
    """
    base = compute_signals(Image.new("RGB", (520, 120), (255, 255, 255)))
    base.update(overrides)
    return base


def test_formula_image_is_classified_as_formula() -> None:
    """
    公式截图必须被判成 formula，而不是"普通图片".

    回归保护，夹具取自真实实测（`_audit_916/audit_parse.py` 生成的公式图，
    PaddleOCR 识别结果 `y= (a+b)/(c-d) * 100%= 12.5`）。

    旧规则下这张图 math_ratio 只算出 0.174 —— 因为 ``_MATH_CHARS`` 里
    **没有 ASCII 括号、斜杠、星号、百分号**，23 个非空白字符里只有 2 个 "="
    和 1 个 "+"、1 个 "-" 被计入。阈值是 0.18，于是差 0.006 被判成普通图片，
    整张图退化成"整图 OCR"，公式引擎（在有的环境下）根本没被调用。

    这里同时锁住另一侧：中文正文即使括号多，也不能被误判成公式。
    """
    from types import SimpleNamespace

    from app.services.image_understanding.classifier import (
        IMAGE_TYPE_FORMULA,
        _decide,
        _text_signals,
    )

    ocr_lines = [SimpleNamespace(text="y= (a+b)/(c-d) * 100%= 12.5")]
    stats = _text_signals(ocr_lines)
    assert stats["math_ratio"] >= 0.18, (
        f"公式的符号占比必须过阈值，实际 {stats['math_ratio']}（旧规则是 0.174）"
    )
    assert stats["cjk_ratio"] == 0.0

    settings = _SETTINGS
    image_type, confidence, reason = _decide(_blank_signals(**stats), settings)
    assert image_type == IMAGE_TYPE_FORMULA, (
        f"公式图应判成 formula，实际 {image_type}（reason={reason}）"
    )
    assert 0.0 < confidence <= 1.0

    # ── 反向：中文正文（括号/百分号很多）不能进公式分支 ──────────────────────
    # 这条反向用例必须**有说服力**：除 cjk 护栏之外的公式条件要全部成立，
    # 否则 math_ratio 自己就过不了阈值，测试是空转 —— 把 cjk 护栏删掉也照样
    # 通过。下面显式锁住"其它条件都成立"，让这道用例真的在测那道护栏。
    prose = [SimpleNamespace(text="(见附录A)占比30%")]
    prose_stats = _text_signals(prose)
    assert prose_stats["math_ratio"] >= 0.18, (
        f"反向夹具的符号占比要过阈值（否则不是有效用例），实际 {prose_stats['math_ratio']}"
    )
    assert prose_stats["cjk_ratio"] > 0.10, "中文正文的 cjk_ratio 应当很高"
    assert prose_stats["code_score"] < 0.5, "中文正文不该像代码"
    prose_type, _, prose_reason = _decide(_blank_signals(**prose_stats), settings)
    assert prose_type != IMAGE_TYPE_FORMULA, (
        f"中文正文不得被判成公式（reason={prose_reason}）"
    )
    print("  ok test_formula_image_is_classified_as_formula")


def test_formula_signals_default_to_non_formula() -> None:
    """
    调用方没提供 cjk_ratio 时（老信号字典）必须保守判成**非公式**.

    `_decide` 的 cjk 护栏缺省值是 1.0 —— 宁可漏判公式，也不要把一张不知底细的
    图丢给公式引擎（真装上了 PP-FormulaNet 时，误判会产出一段假 LaTeX）。

    做成 A/B 对照：同一份信号，**只有** cjk_ratio 的在/不在之差。补齐时判成
    公式，抽掉时判成非公式 —— 这才证明"缺键 => 保守"是这条规则的功劳，而不是
    恰好被别的条件挡住。
    """
    from app.services.image_understanding.classifier import (
        IMAGE_TYPE_FORMULA,
        _decide,
    )

    settings = _SETTINGS
    with_cjk = _blank_signals(math_ratio=0.50)      # 密度足够，其余信号中性
    assert with_cjk["cjk_ratio"] == 0.0
    assert _decide(with_cjk, settings)[0] == IMAGE_TYPE_FORMULA, (
        "补齐 cjk_ratio 时应当判成公式（否则下面的对照没有意义）"
    )

    without_cjk = dict(with_cjk)
    without_cjk.pop("cjk_ratio")                    # 老信号字典：没有这个键
    image_type, _, reason = _decide(without_cjk, settings)
    assert image_type != IMAGE_TYPE_FORMULA, (
        f"缺 cjk_ratio 时必须保守判非公式，实际 {image_type}（reason={reason}）"
    )
    print("  ok test_formula_signals_default_to_non_formula")


def main() -> None:
    print("图片管线 / Document Agent 单元测试")
    test_image_relative_path()
    test_save_and_resolve()
    test_delete_document_images()
    test_searchable_text()
    test_build_image_chunks()
    test_table_content_type()
    test_point_id_and_payload()
    test_document_agent_generates_docx()
    test_docx_table_and_image_extraction()
    test_document_agent_intent()
    # 三层图片处理（分类 → 分流 → 结构化）
    test_classify_image_types()
    test_dark_node_diagram_not_chart()
    test_formula_image_is_classified_as_formula()
    test_formula_signals_default_to_non_formula()
    test_classification_disabled_falls_back_to_ocr()
    test_table_image_content_type_mapping()
    test_structured_content_shape_and_payload()
    test_recognize_table_with_rules_and_cell_ocr()
    test_recognize_table_without_rescue_degrades()
    test_understanding_defaults_are_backward_compatible()
    test_signals_are_serialisable()
    print("\nAll image-pipeline tests passed.")


if __name__ == "__main__":
    main()
