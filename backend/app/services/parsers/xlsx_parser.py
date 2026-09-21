import io
import openpyxl
from app.services.parsers.base import DocumentParser, ExtractionResult, ExtractedPage
from app.utils.logging import get_logger

logger = get_logger(__name__)

class XlsxParser(DocumentParser):
    def parse(self, content: bytes, filename: str, *, document_id: str | None = None, tenant_id: str | None = None) -> ExtractionResult:
        try:
            # data_only=False：保留公式原文（"=A1+B1"），而不是把无缓存值的
            # 公式单元格变 None。值型单元格照常返回其值；公式单元格返回公式。
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
        except Exception as exc:
            raise ValueError(f"Cannot open XLSX '{filename}': {exc}") from exc

        pages: list[ExtractedPage] = []
        cursor = 0
        total_pages = len(wb.sheetnames)

        for i, sheet_name in enumerate(wb.sheetnames):
            sheet = wb[sheet_name]
            sheet_text = [f"Sheet: {sheet_name}"]

            for row in sheet.iter_rows(values_only=True):
                # 保留占位单元格：None 也占一列（写成空串），保证列对齐。
                # 旧实现把 None 过滤掉，导致 a,None,c 被压成 a | c，后续列左移错位。
                row_values = ["" if cell is None else str(cell) for cell in row]
                if not any(v.strip() for v in row_values):
                    continue
                sheet_text.append(" | ".join(row_values))

            page_text = "\n".join(sheet_text).strip()
            # If a sheet only contains its name but no data, we can still include it, but let's check length
            if len(sheet_text) <= 1:
                continue

            start = cursor
            end = start + len(page_text)

            pages.append(ExtractedPage(
                page_number=i + 1,
                text=page_text,
                char_start=start,
                char_end=end
            ))
            cursor = end + 2

        if not pages:
            raise ValueError(f"XLSX '{filename}' contains no extractable text.")

        full_text = "\n\n".join(p.text for p in pages)

        return ExtractionResult(
            pages=pages,
            full_text=full_text,
            # 只数真正产出的 sheet，而非全部 sheet（被 continue 跳过的空 sheet 不计入）。
            page_count=len(pages),
            char_count=len(full_text),
            file_type="xlsx",
            parser_used="openpyxl",
            ocr_used=False,
            ocr_engine=None,
            extraction_method="native",
            is_tabular=True,
        )
